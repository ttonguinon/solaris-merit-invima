# -*- coding: utf-8 -*-
"""
Genera o actualiza la hoja "Version" (primera hoja) de banco-preguntas.xlsx.
Huella: md5 de 10 caracteres sobre id + a + b + c + respuesta de todas las preguntas
funcionales (activas e inactivas), en el orden del Excel.
Uso: python3 version-banco.py [banco-preguntas.xlsx]
"""
import sys, hashlib, datetime, statistics, string
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ARCH = next((a for a in sys.argv[1:] if a.endswith(".xlsx")), "banco-preguntas.xlsx")
wb = load_workbook(ARCH)
F = wb["Funcionales"]; cols = [c.value for c in F[1]]; ix = {c: i for i, c in enumerate(cols)}
filas = [r for r in F.iter_rows(min_row=2, values_only=True) if str(r[0] or "").strip()]

base = "".join("|".join(str(r[ix[k]] or "").strip() for k in ("id", "a", "b", "c", "respuesta")) + "\n" for r in filas)
huella = hashlib.md5(base.encode("utf-8")).hexdigest()[:10]

act = [r for r in filas if str(r[ix["activa"]]).upper() != "NO"]
rep = [r for r in act if "tipo_v2" in ix and str(r[ix["tipo_v2"]] or "").strip().lower() == "repaso"]
pri = [r for r in act if r not in rep]
C = wb["Comportamentales"]; cc = [c.value for c in C[1]]; jx = {c: i for i, c in enumerate(cc)}
comp = [r for r in C.iter_rows(min_row=2, values_only=True) if str(r[0] or "").strip() and str(r[jx["activa"]]).upper() != "NO"]

def ventaja(r):
    k = str(r[ix["respuesta"]]).strip().lower()
    op = {L: str(r[ix[L]] or "") for L in "abc"}
    return len(op[k]) - statistics.mean(len(v) for L, v in op.items() if L != k)
vent = statistics.mean(ventaja(r) for r in pri)

# versión AAAA.MM.DD-letra: si ya hay una del mismo día con otra huella, avanza la letra
hoy = datetime.date.today().strftime("%Y.%m.%d")
previa, huella_previa = None, None
if "Version" in wb.sheetnames:
    for a, b in wb["Version"].iter_rows(min_row=1, max_col=2, values_only=True):
        if a == "Versión": previa = str(b or "")
        if a == "Huella de contenido": huella_previa = str(b or "")
if previa and previa.startswith(hoy) and huella_previa == huella:
    version = previa
elif previa and previa.startswith(hoy):
    letra = previa.split("-")[-1]
    version = f"{hoy}-{string.ascii_lowercase[min(25, string.ascii_lowercase.index(letra) + 1)]}"
else:
    version = f"{hoy}-a"

if "Version" in wb.sheetnames: del wb["Version"]
V = wb.create_sheet("Version", 0)
datos = [
    ("BANCO DE PREGUNTAS · SOLARIS MERIT", ""),
    ("Versión", version),
    ("Fecha", datetime.date.today().isoformat()),
    ("Huella de contenido", huella),
    ("Cómo se calcula la huella", "md5 de id + a + b + c + respuesta de todas las funcionales, 10 caracteres"),
    ("Funcionales activas (simulacros)", len(pri)),
    ("Funcionales de repaso", len(rep)),
    ("Funcionales inactivas", len(filas) - len(act)),
    ("Comportamentales activas", len(comp)),
    ("Ventaja media de longitud de la correcta", round(vent, 1)),
]
for fila in datos: V.append(list(fila))
V["A1"].font = Font(name="Arial", size=13, bold=True, color="FFFFFF"); V["A1"].fill = PatternFill("solid", fgColor="0E1E40")
for r in range(2, len(datos) + 1):
    V.cell(r, 1).font = Font(name="Arial", size=10, bold=True); V.cell(r, 2).font = Font(name="Arial", size=10)
    V.cell(r, 2).alignment = Alignment(horizontal="left")
V.column_dimensions["A"].width = 42; V.column_dimensions["B"].width = 72
wb.save(ARCH)
print(f"Versión {version} · huella {huella} · {len(pri)} principales · {len(rep)} repaso · {len(comp)} comportamentales · ventaja {vent:+.1f} c")
