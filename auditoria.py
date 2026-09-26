# -*- coding: utf-8 -*-
"""
Auditoría del banco de preguntas · Solaris Merit
Uso:  python3 auditoria.py [banco-preguntas.xlsx] [--eje "texto"] [--desde F400]
Mide forma y contenido de los distractores y lista las alertas por id.
"""
import sys, re, unicodedata, statistics
from collections import Counter
from openpyxl import load_workbook

ARCH = next((a for a in sys.argv[1:] if a.endswith(".xlsx")), "banco-preguntas.xlsx")
FILTRO_EJE = None
if "--eje" in sys.argv: FILTRO_EJE = sys.argv[sys.argv.index("--eje")+1].lower()
DESDE = sys.argv[sys.argv.index("--desde")+1] if "--desde" in sys.argv else None

def norm(s):
    s = unicodedata.normalize("NFD", str(s).lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")
def pal(s): return set(re.findall(r"[a-z]{4,}", norm(s)))

ABSOLUTOS = ["nunca","siempre","ninguna","ningun","en ningun caso","jamas","todas las anteriores",
             "exclusivamente","unicamente","solo si","cualquier caso","imposible","obligatoriamente"]
# lenguaje calificado: matices que suenan prudentes y delatan la clave
CALIFICADO = ["conforme a","de acuerdo con","por escrito","dejando constancia","motivad","segun la norma",
              "sin perjuicio","previa","previo","debidamente","oportunamente","con copia","en los terminos",
              "de manera fundamentada","con los soportes","documentad"]
# conducta cautelosa: avisar, consultar, documentar, escalar
CAUTELA = ["informar","informa","informe","reportar","reporta","documentar","documenta","dejar constancia",
           "consultar","consulta","coordinar","coordina","avisar","avisa","poner en conocimiento",
           "solicitar autorizacion","verificar","verifica","requerir","requiere","advertir","advierte",
           "comunicar al superior","al jefe","al coordinador","por escrito"]

def cuenta(txt, lista):
    t = norm(txt); return sum(t.count(k) for k in lista)

wb = load_workbook(ARCH, data_only=True)
F = wb["Funcionales"]; cols = [c.value for c in F[1]]
ix = {c:i for i,c in enumerate(cols)}
filas = []
for r in F.iter_rows(min_row=2, values_only=True):
    d = {c: (r[ix[c]] if ix[c] < len(r) else "") for c in cols}
    if not str(d.get("id") or "").strip(): continue
    if str(d.get("activa","SI")).upper() == "NO": continue
    if FILTRO_EJE and FILTRO_EJE not in str(d.get("eje","")).lower(): continue
    if DESDE and str(d.get("id","")) < DESDE: continue
    filas.append(d)

alertas = {k: [] for k in ["larga","absoluto","calificado","solape","cautela","clave"]}
mas_larga = absolutos = calif_solo = cautela_solo = 0
solapes = []; difs = []
clave = Counter()

for d in filas:
    op = {"A": str(d["a"]), "B": str(d["b"]), "C": str(d["c"])}
    k = str(d["respuesta"]).strip().upper()
    if k not in op: continue
    clave[k] += 1
    cor = op[k]; mal = [v for kk,v in op.items() if kk != k]
    # 1 · longitud
    if len(cor) > max(len(m) for m in mal):
        mas_larga += 1; alertas["larga"].append(d["id"])
    difs.append((max(len(v) for v in op.values()) - min(len(v) for v in op.values())) / max(1, min(len(v) for v in op.values())) * 100)
    # 2 · absolutos en distractores
    if any(cuenta(m, ABSOLUTOS) for m in mal):
        absolutos += 1; alertas["absoluto"].append(d["id"])
    # 3 · lenguaje calificado solo en la correcta
    if cuenta(cor, CALIFICADO) > 0 and all(cuenta(m, CALIFICADO) == 0 for m in mal):
        calif_solo += 1; alertas["calificado"].append(d["id"])
    # 4 · solapamiento enunciado ↔ correcta
    pe = pal(d["enunciado"])
    sc = len(pe & pal(cor)) / max(1, len(pal(cor)))
    sm = max(len(pe & pal(m)) / max(1, len(pal(m))) for m in mal)
    solapes.append(sc*100)
    if sc > 0.40 and sc > sm:
        alertas["solape"].append(d["id"])
    # 5 · conducta cautelosa: la correcta es la más prudente
    cc = cuenta(cor, CAUTELA); cm = max(cuenta(m, CAUTELA) for m in mal)
    if cc > cm:
        cautela_solo += 1; alertas["cautela"].append(d["id"])

n = len(filas)
def pc(x): return f"{x/max(1,n)*100:.1f} %"
print(f"\n=== FUNCIONALES · {n} preguntas activas" + (f" · eje: {FILTRO_EJE}" if FILTRO_EJE else "") + " ===")
print(f"{'MÉTRICA':52} {'RESULTADO':>10}   META")
print(f"{'Correcta es la opción más larga':52} {pc(mas_larga):>10}   < 50 %")
print(f"{'Absolutos en algún distractor':52} {pc(absolutos):>10}   < 3 %")
print(f"{'Lenguaje calificado SOLO en la correcta':52} {pc(calif_solo):>10}   0 %")
print(f"{'Correcta es la más cautelosa (avisar/documentar)':52} {pc(cautela_solo):>10}   < 35 %")
print(f"{'Solapamiento medio enunciado ↔ correcta':52} {statistics.mean(solapes):9.1f} %   < 40 %")
print(f"{'Preguntas con solape alto y mayor que el distractor':52} {pc(len(alertas['solape'])):>10}   < 10 %")
print(f"{'Diferencia media de longitud entre opciones':52} {statistics.mean(difs):9.1f} %   < 60 %")
print(f"{'Balance de clave A / B / C':52} {clave['A']} / {clave['B']} / {clave['C']}")
casos={}
for d in filas: casos.setdefault(d["caso_id"],d)
largos=[len(str(d["texto_caso"])) for d in casos.values()]
ricos=sum(1 for l in largos if l>=220)
jur=sum(1 for d in casos.values() if any(x in str(d["fuente"]) for x in ["C. de E.","Corte Constitucional","T-","C-"]))
import statistics as st2
print(f"{'Casos con contexto (220 caracteres o más)':52} {ricos/max(1,len(casos))*100:8.1f} %   > 60 %")
print(f"{'Longitud media del caso':52} {st2.mean(largos):9.0f} c   200-400")
import re as _re
_NUM=r"(un|una|dos|tres|cuatro|cinco|seis|siete|ocho|nueve|diez|once|doce|quince|veinte|treinta|cuarenta|sesenta|cien|mil)\b"
_T=r"(d[ií]as?|horas?|semanas?|mes(es)?|años?|lunes|martes|mi[eé]rcoles|jueves|viernes|s[aá]bado|domingo|hoy|ayer|ma[ñn]ana|p\. m\.|a\. m\.|vence|plazo|t[eé]rmino)"
_L=r"(Huila|Neiva|Bogot[áa]|Ibagu[ée]|Cartagena|Buenaventura|Rivera|puerto|aeropuerto|municipio|barrio|plaza|regional|sede)"
_pat=_re.compile(r"\d|%|°C|"+_NUM+"|"+_T+"|"+_L, _re.I)
ancla=sum(1 for d in casos.values() if _pat.search(str(d["texto_caso"])))
print(f"{'Casos con ancla concreta (cifra, tiempo o lugar)':52} {ancla/max(1,len(casos))*100:8.1f} %   > 90 %")
print(f"{'Casos apoyados en jurisprudencia':52} {jur/max(1,len(casos))*100:8.1f} %   informativo ({jur} casos)")

# ---- comportamentales ----
C = wb["Comportamentales"]; cc_ = [c.value for c in C[1]]; jx = {c:i for i,c in enumerate(cc_)}
comp = []
for r in C.iter_rows(min_row=2, values_only=True):
    d = {c: (r[jx[c]] if jx[c] < len(r) else "") for c in cc_}
    if not str(d.get("id") or "").strip(): continue
    if str(d.get("activa","SI")).upper() == "NO": continue
    comp.append(d)
larga3 = enc3 = escala3 = 0; al_comp = []
for d in comp:
    op = [str(d["a"]), str(d["b"]), str(d["c"])]; p = [d["puntos_a"], d["puntos_b"], d["puntos_c"]]
    try: i3 = list(p).index(3)
    except ValueError: continue
    if len(op[i3]) > max(len(o) for j,o in enumerate(op) if j != i3): larga3 += 1; al_comp.append(d["id"])
    if " y " in norm(op[i3]): enc3 += 1
    if cuenta(op[i3], ["al jefe","al coordinador","informa","reporta","documenta","constancia","superior"]) > 0: escala3 += 1
m = len(comp)
def pm(x): return f"{x/max(1,m)*100:.1f} %"
# ---- prueba de pistas: cuánto acierta quien no estudió (empates al azar) ----
def _justo(score):
    acc=0.0; nn=0
    for d in filas:
        op={"A":str(d["a"]),"B":str(d["b"]),"C":str(d["c"])}; kk=str(d["respuesta"]).strip().upper()
        if kk not in op: continue
        nn+=1; sc={L:score(v) for L,v in op.items()}; m=max(sc.values()); top=[L for L in sc if sc[L]==m]
        acc+=(1/len(top)) if kk in top else 0
    return acc/max(1,nn)*100
_pos={"más larga":0,"longitud media":0,"más corta":0}
for d in filas:
    _op={L:str(d[L]) for L in "abc"}; _k=str(d["respuesta"]).strip().lower()
    if _k not in _op: continue
    _o=sorted("abc", key=lambda x: len(_op[x]))
    _pos["más corta"]+= _o[0]==_k; _pos["longitud media"]+= _o[1]==_k; _pos["más larga"]+= _o[2]==_k
print("\n=== REPARTO POR LONGITUD · la correcta debe caer en cada posición cerca del 33 % ===")
for _t,_v in _pos.items(): print(f"{('La correcta es la '+_t):52} {_v/max(1,len(filas))*100:8.1f} %   30-37 %")
print("\n=== PRUEBA DE PISTAS · acierto sin estudiar (azar = 33,3 %) ===")
for _t,_f in [("Marcar la más larga",len),("Marcar la más corta",lambda t:-len(t)),
              ("Marcar la que junta dos acciones",lambda t:1 if re.search(r"\by\b|,",t) else 0),
              ("Marcar la más cautelosa",lambda t:cuenta(t,CAUTELA)),("Marcar la más «jurídica»",lambda t:cuenta(t,CALIFICADO))]:
    print(f"{_t:52} {_justo(_f):8.1f} %   < 38 %")
print(f"\n=== COMPORTAMENTALES · {m} situaciones activas ===")
print(f"{'La de 3 puntos es la más larga':52} {pm(larga3):>10}   < 50 %")
print(f"{'La de 3 puntos encadena acciones con «y»':52} {pm(enc3):>10}   < 40 %")
print(f"{'La de 3 puntos avisa, escala o documenta':52} {pm(escala3):>10}   < 50 %")
import statistics as st
largos=[len(str(d.get("situacion",""))) for d in comp]
cons=Counter(str(d.get("consigna","")).strip() for d in comp)
inv=sum(1 for d in comp if str(d.get("invertida","")).upper()=="SI")
ricas=sum(1 for l in largos if l>=180)
print(f"{'Situaciones con contexto (180 caracteres o más)':52} {pm(ricas):>10}   > 80 %")
print(f"{'Longitud media de la situación':52} {st.mean(largos):9.0f} c   180-320")
print(f"{'Consignas distintas en uso':52} {len(cons):>10}   3 o más")
print(f"{'Consigna más repetida':52} {max(cons.values())/max(1,m)*100:8.1f} %   < 60 %")
print(f"{'Situaciones con consigna invertida':52} {pm(inv):>10}   10-20 %")
_ops=[str(d[L]) for d in comp for L in "abc"]
_pt=sum(1 for t in _ops if t.strip().endswith("."))
_mi=sum(1 for t in _ops if t[:1].islower())
def _norm(t):
    import unicodedata as _u
    z=_u.normalize("NFD",str(t).lower()); return "".join(c for c in z if _u.category(c)!="Mn")
def _cola(t,k):
    p=re.findall(r"[a-z0-9]+", _norm(t)); return " ".join(p[-k:]) if len(p)>=k else None
_c=Counter()
for t in _ops:
    for k in (4,5,6):
        x=_cola(t,k)
        if x: _c[x]+=1
_rep=sum(1 for x,v in _c.items() if v>=6)
print(f"{'Opciones que terminan en punto':52} {_pt:>10}   0")
print(f"{'Opciones que empiezan en minúscula':52} {_mi:>10}   0")
print(f"{'Coletillas repetidas en 6 o más opciones':52} {_rep:>10}   0")

print("\n=== ALERTAS POR ID ===")
for k, tit in [("larga","Correcta más larga"),("absoluto","Absolutos en distractor"),
               ("calificado","Calificado solo en la correcta"),("solape","Solape alto"),
               ("cautela","Correcta más cautelosa")]:
    ids = alertas[k]
    print(f"\n{tit} ({len(ids)}):")
    print("  " + (", ".join(map(str, ids)) if ids else "ninguna"))
print(f"\nComportamentales · la de 3 puntos más larga ({len(al_comp)}):")
print("  " + (", ".join(map(str, al_comp)) if al_comp else "ninguna"))
