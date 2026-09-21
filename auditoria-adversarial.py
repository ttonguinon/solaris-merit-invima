# -*- coding: utf-8 -*-
"""
Auditoría adversarial · Solaris Merit
Un "resolvedor tramposo" que NO lee el caso ni el enunciado: solo mira las tres opciones.
Baraja las opciones antes de cada intento (como la app) y desempata al azar.

Uso:  python3 auditoria-adversarial.py [banco-preguntas.xlsx] [--corridas 200] [--ids]
"""
import sys, re, random, unicodedata, statistics, collections
from openpyxl import load_workbook

ARCH = next((a for a in sys.argv[1:] if a.endswith(".xlsx")), "banco-preguntas.xlsx")
CORRIDAS = int(sys.argv[sys.argv.index("--corridas")+1]) if "--corridas" in sys.argv else 200
VER_IDS = "--ids" in sys.argv
META = 40.0

def norm(s):
    s = unicodedata.normalize("NFD", str(s).lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")

CALIF = ["conforme", "segun", "por escrito", "constancia", "verific", "de acuerdo con", "motivad",
         "sin perjuicio", "previa", "previo", "debidamente", "oportunamente", "en los terminos", "document"]
CAUT  = ["informar", "informa ", "reportar", "reporta ", "documentar", "documenta ", "avisar", "avisa ",
         "al jefe", "al coordinador", "al superior", "poner en conocimiento"]
ABS   = ["nunca", "siempre", "ninguna", "ningun ", "en ningun caso", "jamas", "unicamente",
         "exclusivamente", "solo si", "todas las", "todos los", "cualquier caso"]

def n_calif(t): return sum(norm(t).count(k) for k in CALIF)
def n_caut(t):  return sum(norm(t).count(k) for k in CAUT)
def tiene_abs(t): return any(k in norm(t) for k in ABS)
def n_cadena(t): return len(re.findall(r"\s(y|e)\s", norm(t)))
def cola(t, k):
    p = re.findall(r"[a-z0-9]+", norm(t))
    return " ".join(p[-k:]) if len(p) >= k else None
def fin_punto(t): return str(t).strip().endswith((".", ")"))

# ---------- lectura ----------
wb = load_workbook(ARCH, data_only=True)
F = wb["Funcionales"]; cols = [c.value for c in F[1]]; ix = {c: i for i, c in enumerate(cols)}
items = []
for r in F.iter_rows(min_row=2, values_only=True):
    if not str(r[0] or "").strip(): continue
    if str(r[ix["activa"]]).upper() == "NO": continue
    if "tipo_v2" in ix and str(r[ix["tipo_v2"]] or "").strip().lower() == "repaso": continue
    k = str(r[ix["respuesta"]]).strip().upper()
    if k not in "ABC": continue
    items.append({"id": r[ix["id"]], "op": [str(r[ix["a"]]), str(r[ix["b"]]), str(r[ix["c"]])], "k": "ABC".index(k),
                  "expl": str(r[ix["explicacion"]] or ""), "pq": str(r[ix["porque_no"]] or "") if "porque_no" in ix else ""})

# ---------- coletillas repetidas (H7) ----------
todas = [o for it in items for o in it["op"]]
colas = collections.Counter()
for o in todas:
    for k in (4, 5, 6):
        c = cola(o, k)
        if c: colas[c] += 1
REPETIDAS = {c for c, n in colas.items() if n >= 6}
def tiene_cola(t): return any(cola(t, k) in REPETIDAS for k in (4, 5, 6))

# ---------- heurísticas: cada una devuelve un puntaje por opción (se elige el máximo) ----------
def rango(vals):
    s = sorted(set(vals)); return [s.index(v) / max(1, len(s) - 1) for v in vals]
H = {
 "H1 la más larga":                         lambda ops: [len(o) for o in ops],
 "H2 más lenguaje calificado":              lambda ops: [n_calif(o) for o in ops],
 "H3 sin absolutos, la más larga":          lambda ops: [(-1e6 if tiene_abs(o) else 0) + len(o) for o in ops],
 "H4 menciona informar/reportar/avisar":    lambda ops: [n_caut(o) for o in ops],
 "H5 encadena más acciones (y / e)":        lambda ops: [n_cadena(o) for o in ops],
 "H6 combinación ponderada":                lambda ops: [a + b + c + d + e for a, b, c, d, e in zip(
                                                rango([len(o) for o in ops]), rango([n_calif(o) for o in ops]),
                                                rango([n_caut(o) for o in ops]), rango([n_cadena(o) for o in ops]),
                                                [0 if tiene_abs(o) else 1 for o in ops])],
 "H7 descarta coletillas repetidas":        lambda ops: [0 if tiene_cola(o) else 1 for o in ops],
 "H8 la única que termina en punto o )":    lambda ops: [1 if fin_punto(o) else 0 for o in ops],
}
def aplica(nombre, ops):
    s = H[nombre](ops)
    if nombre.startswith("H8"): return sum(1 for o in ops if fin_punto(o)) == 1
    if nombre.startswith("H7"): return 0 < sum(1 for o in ops if tiene_cola(o)) < 3
    return len(set(s)) > 1

rnd = random.Random(2026)
res = {}
for nombre, f in H.items():
    aciertos = 0.0; ap = 0; por_item = {}; acc_ap = 0.0
    for it in items:
        a_ = aplica(nombre, it["op"]); ap += a_
        acc = 0
        for _ in range(CORRIDAS):
            orden = [0, 1, 2]; rnd.shuffle(orden)
            ops = [it["op"][i] for i in orden]; ok = orden.index(it["k"])
            s = f(ops); m = max(s); top = [i for i in range(3) if s[i] == m]
            acc += 1 if rnd.choice(top) == ok else 0
        por_item[it["id"]] = acc / CORRIDAS
        aciertos += acc / CORRIDAS
        if a_: acc_ap += acc / CORRIDAS
    res[nombre] = (aciertos / len(items) * 100, ap / len(items) * 100, por_item, (acc_ap / ap * 100) if ap else 0.0, ap)

# ---------- informe ----------
print(f"\nAUDITORÍA ADVERSARIAL · {ARCH} · {len(items)} preguntas (sin repaso) · {CORRIDAS} barajados por pregunta")
print(f"Azar = 33,3 %   ·   Meta: ninguna heurística por encima de {META:.0f} %\n")
print(f"{'HEURÍSTICA':40} {'ACIERTO':>8} {'APLICA EN':>10} {'ACIERTO CUANDO APLICA':>22}  ESTADO")
for nombre, (acc, ap, _, cond, nap) in res.items():
    print(f"{nombre:40} {acc:7.1f} % {ap:8.1f} % {cond:10.1f} % ({nap:3d} ítems)   {'OK' if acc <= META else 'SUPERA'}")
print(f"\nColetillas repetidas detectadas (4 a 6 palabras, en 6 o más opciones): {len(REPETIDAS)}")
for c in sorted(REPETIDAS, key=lambda c: -colas[c])[:10]: print(f"   «…{c}» · {colas[c]} opciones")

clave = collections.Counter("ABC"[it["k"]] for it in items)
print(f"\nBalance de la clave en el Excel: A {clave['A']} · B {clave['B']} · C {clave['C']}")
disp = [(max(len(o) for o in it["op"]) - min(len(o) for o in it["op"])) / max(1, min(len(o) for o in it["op"])) * 100 for it in items]
vent = [len(it["op"][it["k"]]) - statistics.mean(len(o) for j, o in enumerate(it["op"]) if j != it["k"]) for it in items]
print(f"Dispersión de longitud (más larga vs más corta): media {statistics.mean(disp):.1f} % · mediana {statistics.median(disp):.1f} %")
print(f"Ventaja media de longitud de la correcta: {statistics.mean(vent):+.1f} caracteres")
letras = re.compile(r"\b[abcABC]\)|opci[oó]n(es)?\s+[abcABC]\b")
cit = [it["id"] for it in items if letras.search(it["expl"] + " " + it["pq"])]
print(f"Justificaciones que citan letras: {len(cit)}" + (f"  {cit[:10]}" if cit else ""))
puntos = sum(1 for o in todas if o.strip().endswith("."))
print(f"Opciones que terminan en punto: {puntos} de {len(todas)}")
if VER_IDS:
    peor = sorted(res["H5 encadena más acciones (y / e)"][2].items(), key=lambda x: -x[1])
    print("\nPreguntas donde H5 acierta siempre:", [i for i, v in peor if v == 1.0][:60])
